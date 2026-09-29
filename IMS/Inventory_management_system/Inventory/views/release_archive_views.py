"""Release Archive: every load that left the store, one row per waybill line,
filtered by year + month/quarter. Read-only, open to every authorised user."""
from io import BytesIO

from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Q
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from django.views.generic import ListView

from Inventory.models import MaterialOrder, MaterialTransport, ReleaseLetter

MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
          'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
# One select covers month, quarter and whole year: '' | 'q1'..'q4' | 'm1'..'m12'.
PERIODS = ([('', 'Whole year')]
           + [(f'q{q}', f'Q{q} ({MONTHS[3*q-3]}–{MONTHS[3*q-1]})') for q in range(1, 5)]
           + [(f'm{m}', MONTHS[m-1]) for m in range(1, 13)])
PAGE_SIZES = [25, 50, 100, 200]


def _period(get):
    """(year, period) from the query string; defaults to the current month."""
    now = timezone.localdate()
    try:
        year = int(get.get('year') or now.year)
    except ValueError:
        year = now.year
    period = get.get('period', f'm{now.month}')
    if period not in dict(PERIODS):
        period = f'm{now.month}'
    return year, period


def filter_releases(get):
    year, period = _period(get)
    qs = (MaterialTransport.objects
          .exclude(status='Awaiting Transporter')   # placeholder: nothing has left yet
          .filter(date_dispatched__year=year)
          .select_related('material_order', 'material_order__unit',
                          'material_order__release_letter', 'transporter', 'site_receipt')
          .order_by('-date_dispatched', '-id'))
    if period.startswith('q'):
        q = int(period[1:])
        qs = qs.filter(date_dispatched__month__gte=3*q-2, date_dispatched__month__lte=3*q)
    elif period.startswith('m'):
        qs = qs.filter(date_dispatched__month=int(period[1:]))
    if get.get('region'):
        qs = qs.filter(material_order__region=get['region'])
    if get.get('district'):
        qs = qs.filter(material_order__district=get['district'])
    s = (get.get('q') or '').strip()
    if s:
        qs = qs.filter(Q(material_order__name__icontains=s) | Q(waybill_number__icontains=s)
                       | Q(driver_name__icontains=s) | Q(transporter__name__icontains=s)
                       | Q(material_order__package_number__icontains=s))
    return qs


def _letters(transports):
    """Release letter per order. The FK is set inconsistently, so fall back to
    the letter sharing the order's request code, or its bulk-batch base code
    (REQ-…-3 → REQ-…). One query for the page."""
    from Inventory.views.order_views import _batch_base
    codes = set()
    for t in transports:
        o = t.material_order
        if not o.release_letter_id and o.request_code:
            codes |= {o.request_code, _batch_base(o.request_code)}
    by_code = {}
    for rl in ReleaseLetter.objects.filter(request_code__in=codes).order_by('upload_time'):
        by_code[rl.request_code] = rl   # latest wins
    return lambda o: (o.release_letter or by_code.get(o.request_code)
                      or by_code.get(_batch_base(o.request_code)))


def _documents(t, rl, internal):
    docs = []
    if rl:
        docs.append(('Release file (letter & memo)', reverse('release_letter_detail', args=[rl.pk])))
        for label, f in (('Signed release letter (scan)', rl.pdf_file),
                         ('Approval memo', rl.memo_pdf), ('Release letter', rl.letter_pdf)):
            if f:
                docs.append((label, f.url))
    sr = getattr(t, 'site_receipt', None)
    if sr:
        for label, f in (('Endorsed waybill (site)', sr.waybill_pdf),
                         ('Site acknowledgement', sr.acknowledgement_sheet),
                         ('Site photos', sr.site_photos)):
            if f:
                docs.append((label, f.url))
        # ponytail: the generated waybill keeps its own access rule (internal
        # staff / owning transporter); everyone else sees the endorsed scan above.
        if internal:
            docs.append(('Waybill (system copy)', reverse('download_waybill_pdf', args=[t.pk])))
    return docs


class ReleaseArchiveView(LoginRequiredMixin, ListView):
    template_name = 'Inventory/release_archive.html'
    context_object_name = 'transports'

    def get_paginate_by(self, queryset):
        try:
            n = int(self.request.GET.get('per_page', 50))
        except ValueError:
            n = 50
        return n if n in PAGE_SIZES else 50

    def get_queryset(self):
        return filter_releases(self.request.GET)

    def get(self, request, *args, **kwargs):
        if request.GET.get('export') == 'xlsx':
            return self._excel()
        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        u = self.request.user
        internal = u.is_superuser or u.is_staff or u.groups.filter(name='Management').exists()
        rows = list(ctx['transports'])
        letter_for = _letters(rows)
        for t in rows:
            t.docs = _documents(t, letter_for(t.material_order), internal)
        year, period = _period(self.request.GET)
        releases = MaterialOrder.objects.filter(request_type='Release')
        params = self.request.GET.copy()
        params.pop('page', None)
        ctx.update(
            rows=rows, year=year, period=period, periods=PERIODS, page_sizes=PAGE_SIZES,
            per_page=self.get_paginate_by(None),
            years=range(timezone.localdate().year, 2022, -1),
            regions=releases.exclude(region__isnull=True).exclude(region='')
                            .values_list('region', flat=True).distinct().order_by('region'),
            districts=releases.exclude(district__isnull=True).exclude(district='')
                              .values_list('district', flat=True).distinct().order_by('district'),
            f_region=self.request.GET.get('region', ''),
            f_district=self.request.GET.get('district', ''),
            q=self.request.GET.get('q', ''),
            querystring=params.urlencode(),
            total=ctx['paginator'].count if ctx.get('paginator') else len(rows),
        )
        return ctx

    def _excel(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Release archive'
        head = ['Date', 'Region', 'District', 'Waybill no.', 'Material', 'Quantity',
                'Unit', 'Contract package', 'Driver', 'Transporter']
        ws.append(head)
        for c in range(1, len(head) + 1):
            ws.cell(1, c).font = Font(bold=True, color='FFFFFF')
            ws.cell(1, c).fill = PatternFill('solid', fgColor='1F4E78')
        for t in filter_releases(self.request.GET).iterator():
            o = t.material_order
            ws.append([
                timezone.localtime(t.date_dispatched).strftime('%Y-%m-%d'),
                o.region or '', o.district or '', t.waybill_number or '', o.name,
                float(t.quantity) if t.quantity is not None else None,
                o.unit.name if o.unit else '', o.package_number or '',
                t.driver_name or '', t.transporter.name if t.transporter else ''])
        for i, w in enumerate([12, 16, 20, 16, 40, 12, 10, 18, 22, 26], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        buf = BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        year, period = _period(self.request.GET)
        resp['Content-Disposition'] = f'attachment; filename="release-archive-{year}{("-" + period) if period else ""}.xlsx"'
        return resp
