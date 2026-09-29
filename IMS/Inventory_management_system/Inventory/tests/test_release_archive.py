"""Release Archive: period filtering, visibility to all users, export."""
from datetime import datetime
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from Inventory.models import MaterialOrder, MaterialTransport, ReleaseLetter, Transporter, Unit


def _at(y, m):
    return timezone.make_aware(datetime(y, m, 15, 10, 0))


class ReleaseArchiveTests(TestCase):
    def setUp(self):
        # Any authorised user — here a consultant, not stores staff.
        self.user = User.objects.create_user('viewer', password='pw')
        self.user.groups.add(Group.objects.get_or_create(name='Consultants')[0])
        unit = Unit.objects.create(name='pcs')
        tr = Transporter.objects.create(name='NF3 Logistics')

        def load(name, when, status='In Transit', region='Ashanti', transporter=tr):
            o = MaterialOrder.objects.create(
                name=name, quantity=Decimal('10'), unit=unit, request_type='Release',
                region=region, district='Ejisu', package_number='PKG-7',
                request_code=f'REQ-ARC-001-{MaterialOrder.objects.count() + 1}')
            t = MaterialTransport.objects.create(
                material_order=o, transporter=transporter, driver_name='Kofi',
                waybill_number='WB-1', quantity=Decimal('10'), status=status)
            MaterialTransport.objects.filter(pk=t.pk).update(date_dispatched=when)
            return t

        load('Pole 9m', _at(2026, 2))
        load('Conductor', _at(2026, 5))
        load('Stay wire', _at(2026, 2), region='Volta')
        load('Placeholder', _at(2026, 2), status='Awaiting Transporter', transporter=None)
        # Created after the orders, so their FK stays empty and the archive
        # must find the letter through the bulk batch's base request code.
        self.letter = ReleaseLetter.objects.create(request_code='REQ-ARC-001', title='Arc')
        self.client.force_login(self.user)

    def _get(self, **params):
        return self.client.get(reverse('release_archive'), params)

    def _names(self, **params):
        return {t.material_order.name for t in self._get(**params).context['rows']}

    def test_month_filter(self):
        self.assertEqual(self._names(year=2026, period='m2'), {'Pole 9m', 'Stay wire'})

    def test_quarter_and_year(self):
        self.assertNotIn('Conductor', self._names(year=2026, period='q1'))
        self.assertEqual(self._names(year=2026, period='q2'), {'Conductor'})
        self.assertEqual(self._get(year=2026, period='').context['total'], 3)

    def test_placeholders_that_never_left_are_excluded(self):
        self.assertNotIn('Placeholder', self._names(year=2026, period=''))

    def test_region_filter(self):
        self.assertEqual(self._names(year=2026, period='', region='Volta'), {'Stay wire'})

    def test_release_letter_found_by_request_code(self):
        resp = self._get(year=2026, period='m2')
        self.assertContains(resp, reverse('release_letter_detail', args=[self.letter.pk]))

    def test_excel_export_respects_filters(self):
        from io import BytesIO
        from openpyxl import load_workbook
        resp = self._get(year=2026, period='m2', export='xlsx')
        rows = list(load_workbook(BytesIO(resp.content)).active.values)
        self.assertEqual(len(rows), 3)          # header + 2 February loads
        self.assertEqual({r[4] for r in rows[1:]}, {'Pole 9m', 'Stay wire'})

    def test_bad_params_fall_back_safely(self):
        self.assertEqual(self._get(year='abc', period='zz', per_page='9999').status_code, 200)
